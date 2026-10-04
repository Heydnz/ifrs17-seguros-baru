"""Builds excel/Revision_VT_poliza_modelo.xlsx: the single-policy twin of notebook 08 (live formulas).

Python reference values come from the notebook's own functions (its code cells are executed
up to the projection), so the Conciliacion sheet compares Excel against the notebook.
"""
import json
import math
import os
from pathlib import Path

import nbformat
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel" / "Revision_VT_poliza_modelo.xlsx"
POLICY = "VT2025-05533"

# ---------------------------------------------------------------- Python reference (notebook 08)
nb = nbformat.read(ROOT / "notebooks/08_revision_independiente_vida_temporal.ipynb", 4)
ns = {}
os.chdir(ROOT / "notebooks")
for cell in nb.cells:
    if cell.cell_type != "code":
        continue
    exec(cell.source, ns)
    if "def project(" in cell.source:
        break
project, DF_LOCKED, c25 = ns["project"], ns["DF_LOCKED"], ns["c25"]
assumptions, cso = ns["assumptions"], ns["cso"]
pol = c25[c25["policy_id"] == POLICY]
p = pol.iloc[0]
b = project(pol, DF_LOCKED)
scen = {"Mort": project(pol, DF_LOCKED, mort=1.15), "CadAlza": project(pol, DF_LOCKED, lapse=1.5),
        "CadBaja": project(pol, DF_LOCKED, lapse=0.5), "Gastos": project(pol, DF_LOCKED, expense=1.10, infl_add=0.01)}
fcf = {k: float(v.fcf[0]) for k, v in scen.items()}
fcf["Base"] = float(b.fcf[0])
dm = max(fcf["Mort"] - fcf["Base"], 0)
dl = max(fcf["CadAlza"] - fcf["Base"], fcf["CadBaja"] - fcf["Base"], 0)
de = max(fcf["Gastos"] - fcf["Base"], 0)
cap0 = math.sqrt(dm**2 + dl**2 + de**2 + 2 * 0.25 * dm * de + 2 * 0.5 * dl * de)
cu = b.cu[0]
duration = float((cu * DF_LOCKED[1:21]).sum() / cu[0])
ra = 0.06 * cap0 * duration
csm = max(0.0, -(fcf["Base"] + ra))
fwd0 = DF_LOCKED[0] / DF_LOCKED[1] - 1
rel1 = csm * (1 + fwd0) * cu[0] / cu.sum()
close1 = csm * (1 + fwd0) - rel1
t = np.arange(20)
vp_year = ((b.commissions + b.underwriting + b.maintenance - b.premiums)[0] * DF_LOCKED[:20]
           + (b.claims + b.claim_handling)[0] * DF_LOCKED[1:21])
pv_in = float(b.pv_in[0])
pv_out = float(b.pv_out[0])

# ---------------------------------------------------------------- styles
ARIAL = "Arial"
F = lambda **k: Font(name=ARIAL, **({"size": 10} | k))
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
USD, USD2, PCT, PCT2, Q, DF = '#,##0;(#,##0);"-"', '#,##0.00;(#,##0.00);"-"', '0.0%', '0.00%', '0.000000', '0.000000'

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F(bold=True, size=13, color="1F4E79")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F(italic=True, color="595959")


def header(ws, row, labels, col=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def put(ws, ref, value, color=BLACK, fmt=None, bold=False, fill=None):
    c = ws[ref]
    c.value = value
    c.font = F(color=color, bold=bold)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def section(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = F(bold=True, color="1F4E79")
    c.fill = SEC_FILL


# ---------------------------------------------------------------- Leeme
ws = wb.active
ws.title = "Leeme"
title(ws, "Revisión independiente de vida temporal: modelo de una póliza",
      "Seguros Barú, S.A. (ficticia) · reconocimiento inicial al 1-ene-2025 · gemelo en Excel del notebook 08")
lines = [
    ("Qué es", "La póliza de la sección 3.1 del notebook 08, calculada con fórmulas vivas: proyección de flujos, "
               "valor presente de los flujos de cumplimiento, ajuste por riesgo por costo de capital, CSM y su amortización."),
    ("Para qué", "Mostrar que el mismo cálculo se puede revisar celda por celda en Excel y automatizar en Python, "
                 "y que ambos dan el mismo resultado (hoja Conciliacion)."),
    ("Hojas", "Supuestos: datos de la póliza y supuestos de la entidad · Tablas: CSO 1980 y curva bloqueada · "
              "Proy_Base y cuatro escenarios de choque con la misma estructura · RA_CSM: capital, ajuste por riesgo, "
              "sensibilidad a la tasa, CSM y amortización · Conciliacion: Excel contra el notebook."),
    ("Cómo usarlo", "Cambia solo las celdas en azul de Supuestos (por ejemplo, la tasa de costo de capital, resaltada en amarillo). "
                    "Todo lo demás se recalcula. La conciliación deja de cuadrar si cambias un supuesto, porque los valores "
                    "de Python corresponden a los supuestos de la entidad."),
    ("Colores", "Azul = dato o supuesto escrito a mano · Negro = fórmula · Verde = vínculo a otra hoja · "
                "Amarillo = palanca clave para pruebas de sensibilidad."),
    ("Ajuste por riesgo", "Aquí es el de una póliza sola, sin diversificación. En el notebook el capital se diversifica "
                          "en toda la cohorte y se asigna a cada celda (B88(a)), así que las cifras no son comparables una a una."),
    ("Convenciones", "Pasos anuales. Primas, comisiones y gastos al inicio del año; muertes y gastos de siniestros al final. "
                     "Flujos de cumplimiento = salidas − entradas (negativo = entrada neta)."),
    ("Fuentes", "IFRS 17 (párrafos citados en cada bloque) · SOA, 1980 CSO, tablas 42 y 36 · U.S. Treasury, curva HQM "
                "(datos del repositorio) · Reglamento Delegado (UE) 2015/35, art. 39, 137, 140 y 142 (choques de referencia)."),
    ("Aviso", "Proyecto personal con fines demostrativos. Datos sintéticos de una aseguradora ficticia. No es una opinión "
              "de auditoría ni representa la metodología, las herramientas o la opinión de ningún empleador o cliente."),
]
for k, (a, txt) in enumerate(lines, start=4):
    put(ws, f"A{k}", a, bold=True)
    c = put(ws, f"B{k}", txt)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws[f"A{k}"].alignment = Alignment(vertical="top")
    ws.row_dimensions[k].height = 42
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 110

# ---------------------------------------------------------------- Supuestos
ws = wb.create_sheet("Supuestos")
title(ws, "Supuestos", "Datos de la póliza (censo de la entidad) y documento de supuestos de la entidad")
header(ws, 3, ["Concepto", "Valor", "Fuente o nota"])
rows = [
    ("sec", "Póliza"),
    ("Identificador", POLICY, None, None, "census.csv"),
    ("Sexo", p["sex"], "Sexo", None, "census.csv (M/F)"),
    ("Edad de emisión", int(p["issue_age"]), "EdadEmision", "0", "census.csv"),
    ("Plazo (años)", int(p["term"]), "Plazo", "0", "census.csv; límite del contrato = plazo completo (par. 34)"),
    ("Suma asegurada (USD)", float(p["sum_insured"]), "SumaAsegurada", USD, "census.csv"),
    ("Prima anual nivelada (USD)", float(p["annual_premium"]), "PrimaAnual", USD2, "census.csv"),
    ("Canal", p["channel"], None, None, "census.csv"),
    ("Comisión del primer año", assumptions["commissions"][p["channel"]]["first_year"], "ComisionAnio1", PCT, "assumptions.json"),
    ("Comisión de renovación", assumptions["commissions"][p["channel"]]["renewal"], "ComisionRenov", PCT, "assumptions.json"),
    ("sec", "Mejor estimación (par. 33, B37–B41)"),
    ("Factor sobre CSO 1980, edades 25–39", 0.70, "Fact_25_39", PCT, "assumptions.json; edad alcanzada"),
    ("Factor sobre CSO 1980, edades 40–49", 0.75, "Fact_40_49", PCT, "assumptions.json"),
    ("Factor sobre CSO 1980, edades 50+", 0.95, "Fact_50mas", PCT, "assumptions.json"),
    ("Mejora anual de mortalidad", 0.01, "Mejora", PCT, "assumptions.json; desde 2025 (B58)"),
    ("Caducidad, año póliza 1", 0.08, "Cad_Anio1", PCT, "assumptions.json"),
    ("Caducidad, año póliza 2", 0.06, "Cad_Anio2", PCT, "assumptions.json"),
    ("Caducidad, año póliza 3 en adelante", 0.04, "Cad_Anio3", PCT, "assumptions.json"),
    ("Costo de suscripción por póliza nueva (USD)", 150.0, "CostoSuscripcion", USD2, "assumptions.json; flujo de adquisición"),
    ("Mantenimiento por póliza y año (USD)", 60.0, "CostoMant", USD2, "assumptions.json (B65(h))"),
    ("Costo por siniestro (USD)", 200.0, "CostoSiniestro", USD2, "assumptions.json (B65(f))"),
    ("Inflación de gastos", 0.025, "Inflacion", PCT, "assumptions.json"),
    ("sec", "Ajuste por riesgo: costo de capital (par. 37, B86–B92)"),
    ("Tasa de costo de capital", 0.06, "TasaCoC", PCT2, "Política de la entidad, tomada de Solvencia II (Reg. Delegado 2015/35, art. 39). Ver sensibilidad en RA_CSM."),
    ("Choque de mortalidad", 0.15, "ChoqueMort", PCT, "Referencia: Reg. Delegado 2015/35, art. 137"),
    ("Choque de caducidad (+/−)", 0.50, "ChoqueCad", PCT, "Referencia: art. 142; se toma el peor sentido"),
    ("Choque de gastos", 0.10, "ChoqueGastos", PCT, "Referencia: art. 140"),
    ("Choque de inflación de gastos (puntos)", 0.01, "ChoqueInfl", PCT, "Referencia: art. 140"),
    ("Correlación mortalidad–caducidad", 0.0, "Rho_MC", "0.00", "Referencia: art. 136 (matriz de vida)"),
    ("Correlación mortalidad–gastos", 0.25, "Rho_MG", "0.00", "Referencia: art. 136"),
    ("Correlación caducidad–gastos", 0.50, "Rho_CG", "0.00", "Referencia: art. 136"),
    ("Percentil del capital (a un año)", 0.995, "PercentilCapital", "0.0%", "VaR al 99,5%"),
    ("z del percentil", "=NORMSINV(PercentilCapital)", "Z_Capital", "0.0000", "Fórmula"),
]
r = 4
for row in rows:
    if row[0] == "sec":
        section(ws, f"A{r}", row[1])
        ws.merge_cells(f"A{r}:C{r}")
        r += 1
        continue
    label, val, nm, fmt, note = row
    put(ws, f"A{r}", label)
    is_formula = isinstance(val, str) and val.startswith("=")
    c = put(ws, f"B{r}", val, color=BLACK if is_formula else BLUE, fmt=fmt,
            fill=KEY_FILL if nm == "TasaCoC" else None)
    c.alignment = Alignment(horizontal="right")
    put(ws, f"C{r}", note).font = F(italic=True, color="595959")
    if nm:
        name(nm, f"Supuestos!$B${r}")
    r += 1
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 95
ws.freeze_panes = "A4"

# ---------------------------------------------------------------- Tablas
ws = wb.create_sheet("Tablas")
title(ws, "Tablas públicas", "CSO 1980, edad al cumpleaños más cercano (SOA, tablas 42 y 36) · curva HQM de diciembre de 2024 (bloqueada)")
header(ws, 3, ["Edad", "qx hombres", "qx mujeres"])
for age in range(100):
    put(ws, f"A{4 + age}", age, fmt="0")
    put(ws, f"B{4 + age}", float(cso.loc[age, "qx_male"]), BLUE, "0.00000")
    put(ws, f"C{4 + age}", float(cso.loc[age, "qx_female"]), BLUE, "0.00000")
name("CSO_H", "Tablas!$B$4:$B$103")
name("CSO_M", "Tablas!$C$4:$C$103")
header(ws, 3, ["t (años)", "Factor de descuento bloqueado", "Tasa spot anual"], col=5)
for k in range(22):
    put(ws, f"E{4 + k}", k, fmt="0")
    put(ws, f"F{4 + k}", float(DF_LOCKED[k]), BLUE, DF)
    put(ws, f"G{4 + k}", f"=IF(E{4 + k}=0,\"\",F{4 + k}^(-1/E{4 + k})-1)", fmt=PCT2)
name("FD_Bloqueado", "Tablas!$F$4:$F$25")
ws["E27"] = "Fuente: data/processed/discount_curves_monthly_2024-12.csv, columna 'high' (curva de la política), meses 12, 24, …"
ws["E27"].font = F(italic=True, color="595959")
ws["A105"] = "Fuente: data/raw/mortality/cso1980_anb.csv (mort.soa.org, tablas 42 y 36)."
ws["A105"].font = F(italic=True, color="595959")
for col, w in zip("ABCDEFG", [8, 13, 13, 3, 10, 18, 14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A4"

# ---------------------------------------------------------------- projection sheets
COLS = ["Año póliza", "t", "Edad", "Vigente (par. 34)", "q CSO 1980", "Factor BE", "Mejora",
        "q mejor estimación", "Caducidad", "En vigor al inicio", "Primas", "Comisiones", "Suscripción",
        "Mantenimiento", "Siniestros", "Gasto de siniestros", "FD inicio", "FD fin",
        "VP neto (salidas − entradas)", "Unidades de cobertura (B119)"]
SCEN = {
    "Proy_Base": ("Escenario base (mejor estimación)", "1", "1", "1", "0"),
    "Proy_Mort": ("Choque de mortalidad", "=1+ChoqueMort", "1", "1", "0"),
    "Proy_CadAlza": ("Choque de caducidad al alza", "1", "=1+ChoqueCad", "1", "0"),
    "Proy_CadBaja": ("Choque de caducidad a la baja", "1", "=1-ChoqueCad", "1", "0"),
    "Proy_Gastos": ("Choque de gastos e inflación", "1", "1", "=1+ChoqueGastos", "=ChoqueInfl"),
}
FIRST, LAST = 9, 28
for sname, (desc, m_mort, m_lapse, m_exp, infl) in SCEN.items():
    ws = wb.create_sheet(sname)
    title(ws, f"Proyección de flujos de cumplimiento: {desc.lower()}",
          "Misma estructura en las cinco hojas de proyección; solo cambian los multiplicadores de las celdas E3 a E6.")
    for k, (lab, val) in enumerate([("Multiplicador de mortalidad", m_mort), ("Multiplicador de caducidad", m_lapse),
                                    ("Multiplicador de gastos", m_exp), ("Inflación adicional", infl)]):
        put(ws, f"A{3 + k}", lab)
        ws.merge_cells(f"A{3 + k}:D{3 + k}")
        is_link = val.startswith("=")
        put(ws, f"E{3 + k}", float(val) if not is_link else val, GREEN if is_link else BLUE,
            PCT if k == 3 else "0.00")
    for k, (lab, formula, bold) in enumerate([
            ("VP de primas", f"=SUMPRODUCT(K{FIRST}:K{LAST},Q{FIRST}:Q{LAST})", False),
            ("VP de salidas", f"=SUMPRODUCT(L{FIRST}:L{LAST}+M{FIRST}:M{LAST}+N{FIRST}:N{LAST},Q{FIRST}:Q{LAST})"
                              f"+SUMPRODUCT(O{FIRST}:O{LAST}+P{FIRST}:P{LAST},R{FIRST}:R{LAST})", False),
            ("Flujos de cumplimiento (salidas − entradas)", "=L4-L3", True),
            ("Control: suma de la columna S − L5", f"=ROUND(SUM(S{FIRST}:S{LAST})-L5,6)", False)]):
        put(ws, f"G{3 + k}", lab, bold=bold)
        ws.merge_cells(f"G{3 + k}:K{3 + k}")
        put(ws, f"L{3 + k}", formula, fmt=USD2, bold=bold)
    header(ws, 8, COLS)
    ws.row_dimensions[8].height = 42
    for r in range(FIRST, LAST + 1):
        prev = r - 1
        f = {
            "A": f"=B{r}+1",
            "B": f"=ROW()-{FIRST}",
            "C": f"=EdadEmision+B{r}",
            "D": f"=IF(B{r}<Plazo,1,0)",
            "E": f"=IF(Sexo=\"M\",INDEX(CSO_H,C{r}+1),INDEX(CSO_M,C{r}+1))",
            "F": f"=IF(C{r}<40,Fact_25_39,IF(C{r}<50,Fact_40_49,Fact_50mas))",
            "G": f"=(1-Mejora)^B{r}",
            "H": f"=D{r}*MIN(E{r}*F{r}*G{r}*$E$3,1)",
            "I": f"=D{r}*MIN(IF(B{r}=0,Cad_Anio1,IF(B{r}=1,Cad_Anio2,Cad_Anio3))*$E$4,1)",
            "J": f"=D{r}" if r == FIRST else f"=J{prev}*(1-H{prev})*(1-I{prev})*D{r}",
            "K": f"=J{r}*PrimaAnual",
            "L": f"=K{r}*IF(B{r}=0,ComisionAnio1,ComisionRenov)",
            "M": f"=J{r}*IF(B{r}=0,CostoSuscripcion,0)",
            "N": f"=J{r}*CostoMant*(1+Inflacion+$E$6)^B{r}*$E$5",
            "O": f"=J{r}*H{r}*SumaAsegurada",
            "P": f"=J{r}*H{r}*CostoSiniestro*(1+Inflacion+$E$6)^B{r}*$E$5",
            "Q": f"=INDEX(FD_Bloqueado,B{r}+1)",
            "R": f"=INDEX(FD_Bloqueado,B{r}+2)",
            "S": f"=(L{r}+M{r}+N{r}-K{r})*Q{r}+(O{r}+P{r})*R{r}",
            "T": f"=J{r}*SumaAsegurada",
        }
        fmts = {"A": "0", "B": "0", "C": "0", "D": "0", "E": Q, "F": PCT, "G": "0.0000", "H": Q, "I": PCT,
                "J": "0.000000", "Q": DF, "R": DF, "T": USD}
        for col, formula in f.items():
            put(ws, f"{col}{r}", formula, fmt=fmts.get(col, USD2))
    tr = LAST + 1
    put(ws, f"A{tr}", "Total", bold=True)
    for col in "KLMNOPST":
        put(ws, f"{col}{tr}", f"=SUM({col}{FIRST}:{col}{LAST})", fmt=USD2 if col != "T" else USD, bold=True)
    widths = [9, 5, 6, 10, 10, 9, 8, 11, 10, 11, 11, 11, 11, 13, 11, 12, 10, 10, 14, 16]
    for j, w in enumerate(widths):
        ws.column_dimensions[chr(65 + j)].width = w
    ws.freeze_panes = "C9"

# ---------------------------------------------------------------- RA_CSM
ws = wb.create_sheet("RA_CSM")
title(ws, "Ajuste por riesgo, CSM y amortización", "Póliza sola, sin diversificación (ver Leeme)")
section(ws, "A3", "1. Capital por choque (par. 37, B86)")
header(ws, 4, ["Escenario", "Flujos de cumplimiento", "Diferencia con la base"])
scen_rows = [("Base", "Proy_Base"), ("Mortalidad", "Proy_Mort"), ("Caducidad al alza", "Proy_CadAlza"),
             ("Caducidad a la baja", "Proy_CadBaja"), ("Gastos e inflación", "Proy_Gastos")]
for k, (lab, sh) in enumerate(scen_rows):
    r = 5 + k
    put(ws, f"A{r}", lab)
    put(ws, f"B{r}", f"={sh}!$L$5", GREEN, USD2)
    if k:
        put(ws, f"C{r}", f"=B{r}-$B$5", fmt=USD2)
section(ws, "A11", "2. Capital y ajuste por riesgo (costo de capital)")
items = [
    (12, "Capital de mortalidad", "=MAX(C6,0)", USD2),
    (13, "Capital de caducidad (peor sentido)", "=MAX(C7,C8,0)", USD2),
    (14, "Capital de gastos", "=MAX(C9,0)", USD2),
    (15, "Capital diversificado al inicio", "=SQRT(B12^2+B13^2+B14^2+2*Rho_MC*B12*B13+2*Rho_MG*B12*B14+2*Rho_CG*B13*B14)", USD2),
    (16, "Duración del capital: Σ UC(t)·FD(t+1) / UC(0)", "=SUMPRODUCT(Proy_Base!T9:T28,Proy_Base!R9:R28)/Proy_Base!T9", "0.0000"),
    (17, "Ajuste por riesgo = tasa × capital × duración", "=TasaCoC*B15*B16", USD2),
    (18, "Desviación implícita = capital / z", "=B15/Z_Capital", USD2),
    (19, "Nivel de confianza equivalente (par. 119, aprox. normal)", "=NORMSDIST(B17/B18)", PCT),
]
for r, lab, f, fmt in items:
    put(ws, f"A{r}", lab, bold=r in (17, 19))
    put(ws, f"B{r}", f, GREEN if "Proy_Base" in f else BLACK, fmt, bold=r in (17, 19))
ws["C16"] = "La suma asegurada en vigor esperada proyecta el capital (política de la entidad)."
ws["C19"] = "Aproximación: σ se deduce del capital al 99,5% (Reback, SOA Financial Reporter, 2017)."
for ref in ("C16", "C19"):
    ws[ref].font = F(italic=True, color="595959")
section(ws, "A21", "3. Reconocimiento inicial (par. 38 y 47)")
for r, lab, f in [(22, "Flujos de cumplimiento", "=B5"), (23, "Ajuste por riesgo", "=B17"),
                  (24, "CSM", "=MAX(0,-(B22+B23))"), (25, "Componente de pérdida", "=MAX(0,B22+B23)")]:
    put(ws, f"A{r}", lab, bold=r == 24)
    put(ws, f"B{r}", f, fmt=USD2, bold=r == 24)
section(ws, "A27", "4. Sensibilidad a la tasa de costo de capital")
header(ws, 28, ["Tasa", "Ajuste por riesgo", "Nivel de confianza", "CSM"])
for k, rate in enumerate([0.04, 0.0475, 0.05, 0.06, 0.08]):
    r = 29 + k
    put(ws, f"A{r}", rate, BLUE, PCT2)
    put(ws, f"B{r}", f"=A{r}*$B$15*$B$16", fmt=USD2)
    put(ws, f"C{r}", f"=NORMSDIST(B{r}/$B$18)", fmt=PCT)
    put(ws, f"D{r}", f"=MAX(0,-($B$22+B{r}))", fmt=USD2)
notes = ["4% = Solvencia UK · 4,75% = Solvencia II tras la Directiva (UE) 2025/2 · 6% = Solvencia II original.",
         "Con esta aproximación, nivel de confianza = NORMSDIST(tasa × duración × z): no depende del tamaño del capital.",
         "Referencia de mercado (EY, cierres 2023, 45 aseguradoras): nivel más común 70%–79%; tasas de costo de capital reveladas 4%–6%."]
for k, txt in enumerate(notes):
    ws[f"A{35 + k}"] = txt
    ws[f"A{35 + k}"].font = F(italic=True, color="595959")
section(ws, "A39", "5. Amortización esperada del CSM (par. 44, B72(b), B119)")
header(ws, 40, ["Año póliza", "Saldo inicial", "Tasa forward bloqueada", "Intereses", "Unidades de cobertura",
                "Unidades restantes", "% liberado", "Liberación", "Saldo final"])
S0, S1 = 41, 60
for k in range(20):
    r = S0 + k
    put(ws, f"A{r}", f"=ROW()-{S0 - 1}", fmt="0")
    put(ws, f"B{r}", "=B24" if k == 0 else f"=I{r - 1}", fmt=USD2)
    put(ws, f"C{r}", f"=INDEX(FD_Bloqueado,A{r})/INDEX(FD_Bloqueado,A{r}+1)-1", fmt=PCT2)
    put(ws, f"D{r}", f"=B{r}*C{r}", fmt=USD2)
    put(ws, f"E{r}", f"=Proy_Base!T{FIRST + k}", GREEN, USD)
    put(ws, f"F{r}", f"=SUM(E{r}:E${S1})", fmt=USD)
    put(ws, f"G{r}", f"=IF(F{r}>0,E{r}/F{r},0)", fmt=PCT)
    put(ws, f"H{r}", f"=(B{r}+D{r})*G{r}", fmt=USD2)
    put(ws, f"I{r}", f"=B{r}+D{r}-H{r}", fmt=USD2)
put(ws, f"A{S1 + 1}", "Total", bold=True)
put(ws, f"D{S1 + 1}", f"=SUM(D{S0}:D{S1})", fmt=USD2, bold=True)
put(ws, f"H{S1 + 1}", f"=SUM(H{S0}:H{S1})", fmt=USD2, bold=True)
put(ws, f"A{S1 + 2}", "Control: CSM inicial + intereses − liberación − saldo final")
put(ws, f"I{S1 + 2}", f"=ROUND(B24+D{S1 + 1}-H{S1 + 1}-I{S1},6)", fmt=USD2)
chart = LineChart()
chart.title = "Saldo del CSM al final de cada año póliza"
chart.y_axis.title = "USD"
chart.x_axis.title = "Año póliza"
chart.add_data(Reference(ws, min_col=9, min_row=40, max_row=S1), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=S0, max_row=S1))
chart.height, chart.width = 7.5, 16
chart.legend = None
ws.add_chart(chart, "K40")
for col, w in zip("ABCDEFGHI", [52, 16, 14, 13, 14, 14, 11, 13, 14]):
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- Conciliacion
ws = wb.create_sheet("Conciliacion")
title(ws, "Conciliación: Excel contra el notebook 08 (Python)",
      "Los valores de Python vienen de las funciones del notebook (sección 3) aplicadas a esta póliza.")
put(ws, "A3", "Resultado", bold=True)
header(ws, 5, ["Concepto", "Python (notebook 08)", "Excel", "Diferencia", "¿Cuadra?"])
recon = [
    ("VP de primas", pv_in, "=Proy_Base!L3"),
    ("VP de salidas", pv_out, "=Proy_Base!L4"),
    ("Flujos de cumplimiento, base", fcf["Base"], "=Proy_Base!L5"),
    ("Flujos de cumplimiento, choque de mortalidad", fcf["Mort"], "=Proy_Mort!L5"),
    ("Flujos de cumplimiento, caducidad al alza", fcf["CadAlza"], "=Proy_CadAlza!L5"),
    ("Flujos de cumplimiento, caducidad a la baja", fcf["CadBaja"], "=Proy_CadBaja!L5"),
    ("Flujos de cumplimiento, gastos e inflación", fcf["Gastos"], "=Proy_Gastos!L5"),
    ("Capital diversificado al inicio", cap0, "=RA_CSM!B15"),
    ("Duración del capital", duration, "=RA_CSM!B16"),
    ("Ajuste por riesgo (6%)", ra, "=RA_CSM!B17"),
    ("CSM al reconocimiento inicial", csm, "=RA_CSM!B24"),
    ("Liberación del CSM, año 1", rel1, "=RA_CSM!H41"),
    ("CSM al final del año 1", close1, "=RA_CSM!I41"),
] + [(f"VP neto, año póliza {k + 1}", float(vp_year[k]), f"=Proy_Base!S{FIRST + k}") for k in range(20)]
for k, (lab, py, xl) in enumerate(recon):
    r = 6 + k
    put(ws, f"A{r}", lab)
    put(ws, f"B{r}", round(float(py), 6), BLUE, "#,##0.0000")
    put(ws, f"C{r}", xl, GREEN, "#,##0.0000")
    put(ws, f"D{r}", f"=C{r}-B{r}", fmt="#,##0.000000")
    put(ws, f"E{r}", f"=IF(ABS(D{r})<0.005,\"✔\",\"✘\")")
    ws[f"E{r}"].alignment = Alignment(horizontal="center")
last = 6 + len(recon) - 1
put(ws, "B3", f"=IF(COUNTIF(E6:E{last},\"✘\")=0,\"Todo cuadra (tolerancia USD 0,005)\",\"Revisar diferencias\")", bold=True)
ws["B6"].comment = Comment("Valores exportados de las funciones del notebook 08 con los supuestos de la entidad. "
                           "Si cambias un supuesto en la hoja Supuestos, la conciliación deja de cuadrar a propósito.", "Revisión")
for col, w in zip("ABCDE", [46, 20, 20, 16, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A6"

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("saved", OUT, "| python:", json.dumps({"fcf": round(fcf["Base"], 4), "cap0": round(cap0, 4), "ra": round(ra, 4), "csm": round(csm, 4)}))
