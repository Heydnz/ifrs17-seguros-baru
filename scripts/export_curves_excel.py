"""Export the discount curves of the three closings to Excel, with live formulas.

    python scripts/export_curves_excel.py

Writes excel/Curvas_descuento_Seguros_Baru.xlsx. Treasury inputs are values (blue);
every derived number is an Excel formula that depends on the parameters sheet, so
changing the credit share or a liquidity factor recalculates the whole workbook.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ifrs17.illiquidity import APPLICATION_RATIO, APPLICATION_RATIO_BY_LIQUIDITY, CREDIT_SHARE, CREDIT_SHARE_CDS  # noqa: E402
from ifrs17.treasury_curves import HALF_YEAR_GRID, read_treasury_spot_file  # noqa: E402

RAW = ROOT / "data" / "raw" / "treasury"
CLOSINGS = {"Dic2023": "2023-12", "Dic2024": "2024-12", "Dic2025": "2025-12"}
BUCKETS = [("Alta", "high"), ("Media", "medium"), ("Baja", "low")]

F = "Arial"
BLUE, GREEN, BLACK, GREY = "0000FF", "008000", "000000", "595959"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="DDE6F0")
THIN = Border(bottom=Side(style="thin", color="999999"))
PCT, DF = "0.0000%", "0.000000"
FIRST = 6  # first data row on the half-year sheets
LAST = FIRST + len(HALF_YEAR_GRID) - 1


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)


def header(ws, row, labels, widths=None):
    for j, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=label)
        c.font, c.fill, c.border = font(bold=True), HEAD, THIN
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[row].height = 42


def main() -> None:
    tnc = read_treasury_spot_file(RAW / "tnceom_23_27.xlsx")
    hqm = pd.concat([read_treasury_spot_file(RAW / f) for f in ("hqmeom_19_23.xls", "hqmeom_24_28.xls")])
    hqm = hqm[~hqm.index.duplicated(keep="last")]

    wb = Workbook()
    # ---------------------------------------------------------------- Parameters
    ws = wb.active
    ws.title = "Parametros"
    ws["A1"] = "Seguros Barú — parámetros de la curva de descuento IFRS 17"
    ws["A1"].font = font(bold=True, size=12)
    ws["A2"] = "Celdas amarillas con texto azul: supuestos que se pueden cambiar; todo el libro se recalcula."
    ws["A2"].font = font(GREY, italic=True)
    rows = [("Porcentaje del diferencial HQM − TNC atribuido a crédito", CREDIT_SHARE,
             f"Política elegida: 0% (la curva es la HQM). Opción: {CREDIT_SHARE_CDS:.0%}, con base en CDS "
             "(Longstaff, Mithal y Neis 2005: 51% AAA/AA, 56% A; IFRS 17 B78, B79, B85). Ver hoja Metodo_credito.")]
    rows += [(f"Factor de liquidez — {es}", APPLICATION_RATIO[en],
              {"high": "Vida temporal y su reaseguro (sin valor de rescate).",
               "medium": "Autos, incendio/hogar y TRC.",
               "low": "Vida con ahorro (el asegurado puede rescatar)."}[en]
              + f" Política: 100%. Opción graduada: {APPLICATION_RATIO_BY_LIQUIDITY[en]:.0%}.") for es, en in BUCKETS]
    header(ws, 4, ["Parámetro", "Valor", "Fuente o razón"], [58, 12, 90])
    for i, (label, value, note) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label).font = font()
        c = ws.cell(row=i, column=2, value=value)
        c.font, c.fill, c.number_format = font(BLUE), YELLOW, "0%"
        ws.cell(row=i, column=3, value=note).font = font(GREY)
    ws["A11"] = "Método"
    ws["A11"].font = font(bold=True)
    method = [
        "1. Curva base: TNC del Tesoro de EE. UU., tasas spot de fin de mes, 0,5 a 100 años, capitalización semestral.",
        "2. Prima de iliquidez por plazo = MAX(0; HQM − TNC) × (1 − % crédito) × factor de liquidez (IFRS 17 B79–B80).",
        "3. Tasa semestral = TNC + prima; tasa anual efectiva = (1 + s/2)^2 − 1; factor de descuento = (1 + r)^(−t).",
        "4. Hojas Mensual_: interpolación lineal del logaritmo del factor de descuento entre puntos semestrales; plana por debajo de 0,5 años.",
        "Política elegida: 0% de crédito y factor 100%, así que la curva es exactamente la HQM, porque TNC + (HQM − TNC) = HQM.",
        "Fuente de datos: U.S. Treasury, Office of Economic Policy, archivos tnceom_23_27.xlsx, hqmeom_19_23.xls y hqmeom_24_28.xls.",
    ]
    for i, line in enumerate(method, start=12):
        ws.cell(row=i, column=1, value=line).font = font()
    credit = "Parametros!$B$5"
    factor = {"high": "Parametros!$B$6", "medium": "Parametros!$B$7", "low": "Parametros!$B$8"}

    # ---------------------------------------------------------------- Half-year sheets
    for sheet, month in CLOSINGS.items():
        p = pd.Period(month, "M")
        ws = wb.create_sheet(sheet)
        ws["A1"] = f"Curva de descuento al cierre de {p.strftime('%B %Y')} — malla semestral"
        ws["A1"].font = font(bold=True, size=12)
        ws["A2"] = "TNC y HQM: datos del Tesoro (texto azul). Resto: fórmulas. Tasas como fracción, formato %."
        ws["A2"].font = font(GREY, italic=True)
        labels = ["Plazo (años)", "TNC semestral", "HQM semestral", "Diferencial HQM − TNC (pb)", "Prima neta de crédito"]
        for es, _ in BUCKETS:
            labels += [f"{es}: tasa semestral", f"{es}: tasa anual efectiva", f"{es}: factor de descuento"]
        header(ws, FIRST - 1, labels, [10, 12, 12, 13, 13] + [13] * 9)
        for i, t in enumerate(HALF_YEAR_GRID):
            r = FIRST + i
            ws.cell(row=r, column=1, value=float(t)).font = font()
            for col, src in ((2, tnc), (3, hqm)):
                c = ws.cell(row=r, column=col, value=round(float(src.loc[p, t]), 6))
                c.font, c.number_format = font(BLUE), PCT
            ws.cell(row=r, column=4, value=f"=(C{r}-B{r})*10000").number_format = "0.0"
            c = ws.cell(row=r, column=5, value=f"=MAX(0,C{r}-B{r})*(1-{credit})")
            c.number_format, c.font = PCT, font(GREEN)
            for k, (_, en) in enumerate(BUCKETS):
                cs, ca, cd = (get_column_letter(6 + 3 * k + j) for j in range(3))
                ws[f"{cs}{r}"] = f"=B{r}+E{r}*{factor[en]}"
                ws[f"{ca}{r}"] = f"=(1+{cs}{r}/2)^2-1"
                ws[f"{cd}{r}"] = f"=(1+{ca}{r})^(-A{r})"
                ws[f"{cs}{r}"].font = font(GREEN)
                for ref, fmt in ((cs, PCT), (ca, PCT), (cd, DF)):
                    ws[f"{ref}{r}"].number_format = fmt
        ws.freeze_panes = ws.cell(row=FIRST, column=2)
        ws["D5"].comment = Comment("Positivo en todos los plazos de los tres cierres.", "Claude")

    # ---------------------------------------------------------------- Monthly sheets
    for sheet in CLOSINGS:
        ws = wb.create_sheet(f"Mensual_{sheet}")
        ws["A1"] = f"Curva mensual interpolada — {sheet}"
        ws["A1"].font = font(bold=True, size=12)
        ws["A2"] = ("Interpolación lineal del ln(factor de descuento) entre puntos semestrales de la hoja "
                    f"{sheet}; por debajo de 0,5 años se mantiene la tasa de 0,5 años.")
        ws["A2"].font = font(GREY, italic=True)
        labels = ["Mes", "Plazo (años)", "Semestres completos k"]
        for es, _ in BUCKETS:
            labels += [f"{es}: factor de descuento", f"{es}: tasa anual efectiva", f"{es}: tasa mensual"]
        header(ws, 4, labels, [7, 10, 11] + [13] * 9)
        for m in range(1, 1201):
            r = 4 + m
            ws.cell(row=r, column=1, value=m).font = font()
            ws.cell(row=r, column=2, value=f"=A{r}/12").number_format = "0.0000"
            ws.cell(row=r, column=3, value=f"=INT(A{r}/6)")
            for k in range(3):
                df_src = get_column_letter(8 + 3 * k)  # factor column on the half-year sheet
                rate_src = get_column_letter(7 + 3 * k)
                cdf, can, cmo = (get_column_letter(4 + 3 * k + j) for j in range(3))
                rng = f"{sheet}!${df_src}${FIRST}:${df_src}${LAST}"
                lo = f"INDEX({rng},C{r})"
                hi = f"INDEX({rng},MIN(C{r}+1,{len(HALF_YEAR_GRID)}))"
                ws[f"{cdf}{r}"] = (f"=IF(C{r}=0,(1+{sheet}!${rate_src}${FIRST})^(-B{r}),"
                                   f"EXP(LN({lo})+(LN({hi})-LN({lo}))*(A{r}-6*C{r})/6))")
                ws[f"{can}{r}"] = f"={cdf}{r}^(-1/B{r})-1"
                ws[f"{cmo}{r}"] = f"=(1+{can}{r})^(1/12)-1"
                ws[f"{cdf}{r}"].number_format = DF
                ws[f"{cdf}{r}"].font = font(GREEN)
                ws[f"{can}{r}"].number_format = PCT
                ws[f"{cmo}{r}"].number_format = PCT
        ws.freeze_panes = "D5"

    # ---------------------------------------------------------------- Credit-share methodology
    ws = wb.create_sheet("Metodo_credito")
    ws["A1"] = "Cómo se calcula el porcentaje de crédito (para revisiones futuras)"
    ws["A1"].font = font(bold=True, size=12)
    steps = [
        "IFRS 17 no fija un porcentaje (B78). Pide maximizar datos observables (B78(a)), excluir el riesgo de crédito (B79; B84(b) por analogía) y sugiere los derivados de crédito como referencia (B85).",
        "Paso 1. Participación del incumplimiento según CDS: para cada calificación, prima del CDS ÷ diferencial del bono sobre el Tesoro. Longstaff, Mithal y Neis (2005): 51% en AAA/AA y 56% en A.",
        "Paso 2. Ponderar por la mezcla de calificaciones de la muestra HQM (AAA a A): % crédito = Σ peso × participación. Ejemplo abajo.",
        "Paso 3. Piso de pérdida esperada: % mínimo = probabilidad anual de incumplimiento × severidad ÷ diferencial. Se usa el mayor entre el paso 2 y el paso 3.",
        "Paso 4. Aplicar el mismo % en todos los cierres salvo evidencia nueva (B54–B55) y revisarlo al menos una vez al año.",
        "Paso 5. Revelar el método (par. 117(c)(iii)), las curvas (par. 120) y la sensibilidad del pasivo al % (par. 128).",
    ]
    for i, line in enumerate(steps, start=3):
        ws.cell(row=i, column=1, value=line).font = font()
    ws.column_dimensions["A"].width = 34
    for col in "BCD":
        ws.column_dimensions[col].width = 16
    header(ws, 10, ["Calificación", "Peso en la muestra", "Participación del incumplimiento (CDS)", "Aporte"])
    for i, (rating, weight, share) in enumerate((("AAA/AA", 0.20, 0.51), ("A", 0.80, 0.56)), start=11):
        ws.cell(row=i, column=1, value=rating).font = font()
        for col, val in ((2, weight), (3, share)):
            c = ws.cell(row=i, column=col, value=val)
            c.font, c.fill, c.number_format = font(BLUE), YELLOW, "0%"
        c = ws.cell(row=i, column=4, value=f"=B{i}*C{i}")
        c.number_format = "0.0%"
    ws["A13"], ws["D13"] = "% crédito ponderado (paso 2)", "=SUMPRODUCT(B11:B12,C11:C12)/SUM(B11:B12)"
    ws["A13"].font, ws["D13"].number_format, ws["D13"].font = font(bold=True), "0.0%", font(bold=True)
    for r, (label, val, fmt) in enumerate((("Probabilidad anual de incumplimiento, A", 0.0008, "0.00%"),
                                           ("Severidad (pérdida dado incumplimiento)", 0.60, "0%"),
                                           ("Diferencial HQM − TNC a 10 años, dic-2025", "=Dic2025!D25/10000", "0.00%")), start=15):
        ws.cell(row=r, column=1, value=label).font = font()
        c = ws.cell(row=r, column=4, value=val)
        c.number_format = fmt
        c.font = font(BLUE) if not str(val).startswith("=") else font(GREEN)
        if not str(val).startswith("="):
            c.fill = YELLOW
    ws["A18"], ws["D18"] = "Piso por pérdida esperada (paso 3)", "=D15*D16/D17"
    ws["D18"].number_format = "0.0%"
    ws["A19"], ws["D19"] = "% crédito sugerido = MAX(paso 2; paso 3)", "=MAX(D13,D18)"
    ws["A19"].font, ws["D19"].font, ws["D19"].number_format = font(bold=True), font(bold=True), "0.0%"
    ws["A20"] = ("Valores de ejemplo (amarillo): mezcla de calificaciones supuesta, no publicada por el Tesoro; "
                 "probabilidad y severidad ilustrativas. Actualizarlas con estudios públicos de incumplimiento en cada revisión.")
    ws["A20"].font = font(GREY, italic=True)
    ws["A21"] = "Para aplicar el resultado, copiar D19 en Parametros!B5."
    ws["A21"].font = font(GREY, italic=True)

    # ---------------------------------------------------------------- Summary
    ws = wb.create_sheet("Resumen", 1)
    ws["A1"] = "Tasas spot anuales efectivas a plazos seleccionados"
    ws["A1"].font = font(bold=True, size=12)
    ws["A2"] = "Fórmulas que leen las hojas de cada cierre; cambian si cambian los parámetros."
    ws["A2"].font = font(GREY, italic=True)
    mats = [1, 5, 10, 20, 30, 40, 60, 100]
    header(ws, 4, ["Cierre", "Curva"] + [f"{m} años" for m in mats], [10, 26] + [10] * len(mats))
    r = 5
    for sheet in CLOSINGS:
        curves = [("TNC (Tesoro)", None), ("IFRS 17 — liquidez alta", "G"), ("IFRS 17 — liquidez media", "J"),
                  ("IFRS 17 — liquidez baja", "M"), ("HQM (corporativos)", None)]
        for name, col in curves:
            ws.cell(row=r, column=1, value=sheet).font = font()
            ws.cell(row=r, column=2, value=name).font = font()
            for j, m in enumerate(mats):
                match = f"MATCH({m},{sheet}!$A${FIRST}:$A${LAST},0)"
                if col:
                    f = f"=INDEX({sheet}!${col}${FIRST}:${col}${LAST},{match})"
                else:
                    src = "B" if name.startswith("TNC") else "C"
                    f = f"=(1+INDEX({sheet}!${src}${FIRST}:${src}${LAST},{match})/2)^2-1"
                c = ws.cell(row=r, column=3 + j, value=f)
                c.number_format, c.font = "0.00%", font(GREEN)
            r += 1
        r += 1

    for s in wb.worksheets:
        s.sheet_view.showGridLines = s.title.startswith(("Dic", "Mensual"))
    out = ROOT / "excel" / "Curvas_descuento_Seguros_Baru.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
